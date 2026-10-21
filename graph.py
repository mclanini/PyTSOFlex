from pathlib import Path

import plotly.graph_objects as go
from plotly.io import to_html
import pandas as pd
import webbrowser
import html
import textwrap
import math
from datetime import date

from datetime import datetime

class PlotFlexibilidade:
    """
    Classe responsável por construir as regiões de flexibilidade
    para cada tipo de evento (Tensão, Fluxo, Geração e Divergência).

    A região é construída selecionando, para cada direção (ângulo),
    o primeiro ponto onde o evento ocorreu.
    """

    def __init__(self, df, df_2):
        """
        Inicializa a classe.

        Parameters
        ----------
        df : pandas.DataFrame
            DataFrame contendo todos os resultados da varredura.
        """

        # Cria uma cópia para evitar alterações no DataFrame original.
        self.df = df.copy()
        self.df_2 = df_2.copy()

        # Remove espaços extras dos nomes das colunas.
        self.df.columns = self.df.columns.str.strip()

        # Garante que as colunas numéricas estejam no formato correto.
        self.df["PL"] = pd.to_numeric(self.df["PL"])
        self.df["QL"] = pd.to_numeric(self.df["QL"])
        if "Descrição" not in self.df.columns:
            raise ValueError("O DataFrame precisa conter a coluna 'Descrição'.")

        # Converte o ângulo de "45°" para 45.0.
        self.df["angulo"] = pd.to_numeric(
            self.df["Ângulo"]
                .astype(str)
                .str.replace("°", "", regex=False),
            errors="coerce"
        )

        # Armazena o ponto de operação original.
        self.original = (
            self.df[
                self.df["Característica"] == "Original"
            ]
            .iloc[0]
        )

    def obter_borda(self, evento):
        """
        Obtém os pontos que formam a borda da região de flexibilidade.

        Para cada ângulo, procura o primeiro ponto onde o evento ocorreu.

        Parameters
        ----------
        evento : str
            Pode ser:
                - "Violação de Tensão"
                - "Violação de Fluxo"
                - "Violação de Geração"
                - "Divergência"

        Returns
        -------
        pandas.DataFrame
            DataFrame contendo apenas os pontos que pertencem
            à borda da região.
        """

        borda = []

        # Percorre cada direção estudada.
        for angulo, grupo in (
            self.df[self.df["Característica"].str.contains("Modificado", na=False)]
            .groupby("angulo")
        ):

            # Garante que os pontos estejam na ordem em que
            # foram simulados.
            grupo = grupo.sort_index().reset_index(drop=True)

            # Divergência é identificada quando o fluxo
            # deixa de convergir.
            if evento == "Divergência":

                candidatos = grupo[
                    grupo["Convergiu"] == "Não"
                ]

            # Demais eventos possuem uma coluna própria.
            else:

                candidatos = grupo[
                    grupo[evento] == "Sim"
                ]

            # Caso exista o evento nesta direção,
            # utiliza apenas o primeiro ponto.
            if not candidatos.empty:

                borda.append(candidatos.iloc[0])
        
        # Caso nenhum ponto tenha sido encontrado,
        # retorna um DataFrame vazio.
        if not borda:
            return pd.DataFrame()
        # Organiza os pontos pela ordem angular,
        # formando corretamente o polígono.
        return (
            pd.DataFrame(borda)
            .sort_values("angulo")
            .reset_index(drop=True)
        )

    EVENTOS = ("Violação de Tensão", "Violação de Fluxo",
               "Violação de Geração", "Divergência")
    CORES = {"Violação de Tensão": "#FFFF00",
             "Violação de Fluxo": "#0000FF",
             "Violação de Geração": "#E54CCC",
             "Divergência": "#00FFFF"}

    @staticmethod
    def _eventos_linha(linha):
        eventos = [evento for evento in PlotFlexibilidade.EVENTOS[:-1]
                   if linha.get(evento) == "Sim"]
        if linha.get("Convergiu") == "Não":
            eventos.append("Divergência")
        return eventos

    @staticmethod
    def _descricao_formatada(descricao, largura=48):
        if pd.isna(descricao):
            return "Não informada"
        texto = str(descricao)
        linhas = []
        for paragrafo in texto.splitlines():
            linhas.extend(textwrap.wrap(paragrafo, width=largura,
                                        break_long_words=True,
                                        break_on_hyphens=False) or [""])
        return "<br>".join(html.escape(linha) for linha in linhas)

    @staticmethod
    def _chave_ponto(linha):
        # Coordenadas exatas: não agrupa pontos próximos indevidamente.
        return (float(linha["PL"]), float(linha["QL"]))

    def _registros_compartilhados(self):
        """Mapeia cada coordenada aos eventos e descrições de suas bordas."""
        mapa = {}
        for evento in self.EVENTOS:
            for _, linha in self.obter_borda(evento).iterrows():
                mapa.setdefault(self._chave_ponto(linha), []).append((evento, linha))
        return mapa

    def _hover_compartilhado(self, linha, registros, eventos_extras=()):
        """Exibe todas as causas e elimina descrições duplicadas."""
        eventos = list(dict.fromkeys(
            [evento for evento, _ in registros] + list(eventos_extras)
        ))
        eventos = [e for e in self.EVENTOS if e in eventos]
        # A mesma descrição pode corresponder a mais de um evento.
        descricoes = {}
        for evento, registro in registros:
            valor = registro.get("Descrição")
            descricao = "Não informada" if pd.isna(valor) else str(valor)
            descricoes.setdefault(descricao, [])
            if evento not in descricoes[descricao]:
                descricoes[descricao].append(evento)
        # Eventos que ocorrem no ponto mas não constam das bordas individuais.
        for evento in eventos_extras:
            if not any(evento in associados for associados in descricoes.values()):
                valor = linha.get("Descrição")
                descricao = "Não informada" if pd.isna(valor) else str(valor)
                descricoes.setdefault(descricao, [])
                if evento not in descricoes[descricao]:
                    descricoes[descricao].append(evento)
        cabecalho = (
            f"<b>Ângulo:</b> {linha['angulo']:.0f}°<br>"
            f"<b>P:</b> {linha['PL']:.2f}<br>"
            f"<b>Q:</b> {linha['QL']:.2f}<br>"
            f"<b>Evento Limitante:</b> {html.escape(' e '.join(eventos))}<br>"
        )
        if len(descricoes) == 1:
            descricao = next(iter(descricoes))
            return cabecalho + "<b>Descrição:</b><br>" + self._descricao_formatada(descricao)
        blocos = ["<b>Descrições:</b>"]
        for descricao, associados in descricoes.items():
            rotulo = html.escape(' e '.join(associados))
            blocos.append(f"<b>{rotulo}:</b><br>{self._descricao_formatada(descricao)}")
        return cabecalho + "<br>".join(blocos)

    def obter_borda_flexibilidade(self):
        """Primeiro ponto com qualquer evento, na ordem de simulação."""
        borda = []
        grupos = self.df[self.df["Característica"].str.contains("Modificado", na=False)]
        for angulo, grupo in grupos.groupby("angulo"):
            for _, linha in grupo.sort_index().iterrows():
                eventos = self._eventos_linha(linha)
                if eventos:
                    linha = linha.copy()
                    linha["Evento"] = " e ".join(eventos)
                    borda.append(linha)
                    break
        if not borda:
            return pd.DataFrame()
        return pd.DataFrame(borda).sort_values("angulo").reset_index(drop=True)

    def _hover(self, linha, eventos=None, ponto_original=False):
        descricao = self._descricao_formatada(linha["Descrição"])
        if ponto_original:
            return (f"<b>Ponto de Operação Caso Base</b><br>"
                    f"<b>P:</b> {linha['PL']:.2f}<br>"
                    f"<b>Q:</b> {linha['QL']:.2f}<br>"
                    f"<b>Descrição:</b><br>{descricao}")
        evento_texto = " e ".join(eventos or [])
        return (f"<b>Ângulo:</b> {linha['angulo']:.0f}°<br>"
                f"<b>P:</b> {linha['PL']:.2f}<br>"
                f"<b>Q:</b> {linha['QL']:.2f}<br>"
                f"<b>Evento Limitante:</b> {html.escape(evento_texto)}<br>"
                f"<b>Descrição:</b><br>{descricao}")

    def _adicionar_regiao(self, fig, borda, nome, cor, eventos_compartilhados=None,
                         flexibilidade=False, mostrar_marcadores=True):
        if borda.empty:
            return
        fechada = pd.concat([borda, borda.iloc[[0]]], ignore_index=True)
        textos = []
        for _, linha in borda.iterrows():
            if flexibilidade:
                eventos = self._eventos_linha(linha)
                # Se o mesmo P/Q é uma borda de mais de um evento, exibe todos.
                eventos = [e for e in self.EVENTOS if e in
                           (set(eventos) | (eventos_compartilhados or {}).get(self._chave_ponto(linha), set()))]
            elif eventos_compartilhados is not None:
                eventos = [e for e in self.EVENTOS if e in
                           eventos_compartilhados.get(self._chave_ponto(linha), {nome})]
            else:
                eventos = [nome]
            textos.append(self._hover(linha, eventos))
        textos.append(textos[0])
        fig.add_trace(go.Scatter(
            x=fechada["PL"], y=fechada["QL"],
            mode="lines+markers" if mostrar_marcadores else "lines",
            fill="toself",
            fillcolor="#00FF99" if flexibilidade else cor,
            line=dict(color="#00FF99" if flexibilidade else cor, width=3 if flexibilidade else 2),
            marker=dict(color="black", size=8),
            text=textos, hovertemplate="%{text}<extra></extra>",
            name=nome, hoveron="points" if mostrar_marcadores else "fills",
            hoverinfo="skip" if not mostrar_marcadores else None,
            showlegend=True))

    def _adicionar_operacao(self, fig):
        fig.add_trace(go.Scatter(
            x=[self.original["PL"]], y=[self.original["QL"]],
            mode="markers", marker=dict(symbol="star", size=18, color="black"),
            name="Ponto de Operação Caso Base",
            text=[self._hover(self.original, ponto_original=True)],
            hovertemplate="%{text}<extra></extra>"))

    def _configurar_figura(self, fig, titulo):

        grafico_final = titulo == "Região de Flexibilidade"

        fig.update_layout(
            title=dict(
                text=titulo,
                y=0.98,
                yanchor="top"
            ),
            template="plotly_white",
            height=520,
            autosize=True,
            margin=dict(
                l=40,
                r=20,
                t=85 if grafico_final else 50,
                b=40
            ),
            hovermode="closest",
            hoverdistance=8,
            legend=dict(
                orientation="h",
                xref="container",
                x=0.04,
                xanchor="left",
                yref="container",
                y=0.83 if grafico_final else 0.89,
                yanchor="bottom"
            )
        )

        fig.update_xaxes(
            title="P<sub>load</sub>",
            showgrid=True
        )

        fig.update_yaxes(
            title="Q<sub>load</sub>",
            showgrid=True,
            scaleanchor="x",
            scaleratio=1
        )

        return fig

    def plot(self, evento, cor):
        """Gráfico individual: somente a região do evento e ponto original."""
        borda = self.obter_borda(evento)
        if borda.empty:
            return None
        fig = go.Figure()
        self._adicionar_regiao(fig, borda, evento, cor)
        self._adicionar_operacao(fig)
        return self._configurar_figura(fig, evento)

    def plot_flexibilidade(self):
        """Todas as regiões, com pontos coincidentes consolidados no hover."""
        fig = go.Figure()
        registros = self._registros_compartilhados()
        pontos = {}
        # Desenha primeiro as regiões maiores, para não encobrir as menores.
        # O preenchimento é opaco: não há mistura de cores entre camadas.
        regioes = []
        for evento, cor in self.CORES.items():
            borda = self.obter_borda(evento)
            if borda.empty:
                continue
            x, y = borda["PL"].to_numpy(), borda["QL"].to_numpy()
            area = abs(sum(x[i] * y[(i + 1) % len(x)] -
                           x[(i + 1) % len(x)] * y[i]
                           for i in range(len(x)))) / 2
            regioes.append((area, evento, cor, borda))

        for _, evento, cor, borda in sorted(regioes, key=lambda item: -item[0]):
            self._adicionar_regiao(fig, borda, evento, cor,
                                  mostrar_marcadores=False)
            for _, linha in borda.iterrows():
                pontos.setdefault(self._chave_ponto(linha), linha)

        flex = self.obter_borda_flexibilidade()
        if not flex.empty:
            self._adicionar_regiao(fig, flex, "Região de Flexibilidade", "#00FF99",
                                  flexibilidade=True, mostrar_marcadores=False)
            for _, linha in flex.iterrows():
                pontos[self._chave_ponto(linha)] = linha

        # Um único marcador interativo por coordenada. O ponto de operação
        # é desenhado depois, com prioridade, caso compartilhe coordenadas.
        chave_original = self._chave_ponto(self.original)
        linhas = [linha for chave, linha in pontos.items() if chave != chave_original]
        if linhas:
            textos = [self._hover_compartilhado(
                linha,
                registros.get(self._chave_ponto(linha), []),
                self._eventos_linha(linha)
            ) for linha in linhas]
            fig.add_trace(go.Scatter(
                x=[linha["PL"] for linha in linhas],
                y=[linha["QL"] for linha in linhas],
                mode="markers", marker=dict(color="black", size=9),
                name="Pontos de Violação", showlegend=False,
                text=textos, hovertemplate="%{text}<extra></extra>"))
        self._adicionar_operacao(fig)
        return self._configurar_figura(fig, "Região de Flexibilidade")

    @staticmethod
    def _salvar_excel(df, destino):
        """Exporta os dados para Excel usando pandas e openpyxl."""
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(destino, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Dados", index=False)
            planilha = writer.sheets["Dados"]
            planilha.freeze_panes = "A2"
            planilha.auto_filter.ref = planilha.dimensions
            for celula in planilha[1]:
                celula.fill = PatternFill("solid", fgColor="17365D")
                celula.font = Font(color="FFFFFF", bold=True)
                celula.alignment = Alignment(vertical="center")
            planilha.row_dimensions[1].height = 26
            for indice, coluna in enumerate(df.columns, start=1):
                planilha.column_dimensions[get_column_letter(indice)].width = min(
                    max(len(str(coluna)) + 3, 18), 55
                )

    def gerar_relatorio(self, pasta_saida="Graficos", abrir_html=True):
        from datetime import datetime

        pasta = Path(pasta_saida)
        pasta.mkdir(exist_ok=True)

        cores = self.CORES
        pasta_arquivos = pasta / "arquivos"
        pasta_arquivos.mkdir(exist_ok=True)
        self._salvar_excel(self.df, pasta_arquivos / "resultados_varredura.xlsx")
        self._salvar_excel(self.df_2, pasta_arquivos / "resultados_compilados.xlsx")

        html = f"""
        <!DOCTYPE html>
        <html lang="pt-BR">
        <head>
            <meta charset="UTF-8">
            <title>Regiões de Flexibilidade</title>
            <style>
                * {{
                    box-sizing: border-box;
                }}
                body {{
                    margin: 0;
                    padding: 35px;
                    background: #f3f5f8;
                    font-family: "Segoe UI", Arial, sans-serif;
                    color: #333;
                }}
                .container {{
                    width: min(95vw, 1500px);
                    margin: auto;
                }}
                .plotly-graph-div {{
                    width: 100% !important;
                }}
                h1 {{
                    text-align: center;
                    margin-bottom: 5px;
                }}
                .subtitle {{
                    text-align: center;
                    color: #666;
                    margin-bottom: 35px;
                }}
                .grid {{
                    display: grid;
                    grid-template-columns: repeat(2, minmax(0, 1fr));
                    gap: 25px;
                }}
                .card {{
                    background: white;
                    border-radius: 12px;
                    padding: 20px;
                    box-shadow: 0 4px 15px rgba(0,0,0,.08);
                }}
                .card h2 {{
                    margin-top: 0;
                    text-align: center;
                    margin-bottom: 20px;
                }}
                .downloads-intro {{
                    text-align: center; color: #64748b; margin: 0 0 26px;
                    font-size: 14px; line-height: 1.65;
                }}
                .download-list {{ display: grid; gap: 16px; }}
                .download-item {{
                    display: flex; align-items: center; gap: 15px;
                    padding: 18px; border: 1px solid #e2e8f0;
                    border-radius: 12px; background: #f8fafc;
                    transition: box-shadow .2s, border-color .2s, transform .2s;
                }}
                .download-item:hover {{ border-color: #93c5fd; box-shadow: 0 7px 20px rgba(15,23,42,.07); transform: translateY(-2px); }}
                .file-icon {{
                    width: 48px; height: 48px; flex: 0 0 48px;
                    display: grid; place-items: center; border-radius: 12px;
                    background: #dcfce7; color: #15803d;
                    font-weight: 800; font-size: 15px;
                }}
                .file-meta {{ flex: 1; min-width: 0; }}
                .file-name {{ font-size: 15px; font-weight: 700; color: #0f172a; overflow-wrap: anywhere; }}
                .file-info {{ font-size: 12px; color: #64748b; margin-top: 5px; }}
                .download-link {{
                    text-decoration: none; border-radius: 9px; background: #166534;
                    color: white; padding: 11px 15px; font-weight: 700;
                    font-size: 13px; white-space: nowrap;
                }}
                .download-link:hover {{ background: #14532d; }}
                .download-note {{
                    margin-top: 26px; padding: 15px 18px; border-radius: 10px;
                    background: #eff6ff; color: #475569; font-size: 13px; line-height: 1.6;
                }}
                @media (max-width: 900px) {{ .grid {{ grid-template-columns: 1fr; }} }}
                @media (max-width: 530px) {{ .download-item {{ flex-wrap: wrap; }} .download-link {{ width: 100%; text-align: center; }} }}
                .sem-evento {{
                    height: 480px;
                    display: flex;
                    align-items: center;
                    justify-content: center;
                    color: #777;
                    font-size: 18px;
                }}

            </style>
        </head>
        <body>
            <div class="container">
                <h1>Regiões de Flexibilidade</h1>
                <p class="subtitle">
                    Relatório gerado em
                    {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}
                </p>
                <div class="grid">
        """
        primeiro = True
        graficos = [(evento, self.plot(evento, cor)) for evento, cor in cores.items()]
        graficos.append(("Região de Flexibilidade", self.plot_flexibilidade()))
        for evento, fig in graficos:
            html += f"""
            <div class="card">

                <h2>{evento}</h2>
            """
            if fig is None:
                html += """
                <div class="sem-evento">
                    Nenhuma ocorrência encontrada.
                </div>
                """
            else:
                html += to_html(
                    fig,
                    include_plotlyjs="cdn" if primeiro else False,
                    full_html=False,
                    config={
                        "displaylogo": False,
                        "responsive": True,
                        "scrollZoom": True
                    }
                )
                primeiro = False
            html += """
            </div>
            """
        html += f"""
            <div class="card">
                <h2>Arquivos de Compilação</h2>
                <p class="downloads-intro">
                    Acesse os dados utilizados na elaboração do relatório.<br>
                    Arquivos Excel prontos para análise e armazenamento.
                </p>
                <div class="download-list">
                    <div class="download-item">
                        <div class="file-icon">XLS</div>
                        <div class="file-meta">
                            <div class="file-name">Resultados da Varredura</div>
                            <div class="file-info">resultados_varredura.xlsx · {len(self.df):,} registros</div>
                        </div>
                        <a class="download-link" href="arquivos/resultados_varredura.xlsx" download>↓ Baixar</a>
                    </div>
                    <div class="download-item">
                        <div class="file-icon">XLS</div>
                        <div class="file-meta">
                            <div class="file-name">Resultados Compilados</div>
                            <div class="file-info">resultados_compilados.xlsx · {len(self.df_2):,} registros</div>
                        </div>
                        <a class="download-link" href="arquivos/resultados_compilados.xlsx" download>↓ Baixar</a>
                    </div>
                </div>
                <div class="download-note">
                    Os arquivos contêm os dados completos dos DataFrames originais.
                    Para compartilhar este relatório, mantenha a pasta <b>arquivos</b>
                    junto do <b>index.html</b>.
                </div>
            </div>
        """
        html += """
                </div>
            </div>
        </body>
        </html>
        """
        caminho = pasta / "index.html"
        caminho.write_text(
            html,
            encoding="utf8"
        )
        if abrir_html:
            webbrowser.open(caminho.resolve().as_uri())